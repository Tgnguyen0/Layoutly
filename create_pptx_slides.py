import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()

    # 16:9 widescreen slides
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color palette
    BG_DARK = RGBColor(15, 23, 42)       # #0f172a
    CARD_BG = RGBColor(30, 41, 59)       # #1e293b
    CARD_INNER = RGBColor(15, 23, 42)    # #0f172a
    TEXT_WHITE = RGBColor(248, 250, 252) # #f8fafc
    TEXT_MUTED = RGBColor(148, 163, 184) # #94a3b8
    COLOR_BLUE = RGBColor(56, 189, 248)  # #38bdf8
    COLOR_PURPLE = RGBColor(192, 132, 252)# #c084fc
    COLOR_GREEN = RGBColor(74, 222, 128) # #4ade80
    COLOR_ROSE = RGBColor(251, 113, 133) # #fb7185
    BORDER_COLOR = RGBColor(51, 65, 85)  # #334155

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text):
        # Tag
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = tag_text.upper()
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_BLUE
        p.font.name = 'Arial'

        # Title
        tx_box2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf2 = tx_box2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(24)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE
        p2.font.name = 'Arial'

    # ----------------------------------------------------
    # SLIDE 1: Title Slide
    # ----------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    add_bg(slide1)

    # Title Box
    tb = slide1.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11.333), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "BÁO CÁO NGHĨEN CỨU CÁC GIẢI PHÁP FIGMA-TO-CODE MÃ NGUỒN MỞ"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = "Phân Tích Kiến Trúc, So Sánh Ưu/Nhược Điểm & Đề Xuất Cho Đề Tài LAYOUTLY"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(20)

    p3 = tf.add_paragraph()
    p3.text = "Đánh giá thực nghiệm 4 dự án GitHub (FigmaToCode, Grida, Figma to React, Figma Context MCP)"
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_MUTED
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(15)

    p4 = tf.add_paragraph()
    p4.text = "\nSinh viên thực hiện đề tài Layoutly | Đồ án NCKH / Đồ án Tốt nghiệp"
    p4.font.size = Pt(14)
    p4.font.color.rgb = COLOR_PURPLE
    p4.alignment = PP_ALIGN.CENTER

    # ----------------------------------------------------
    # SLIDE 2: Tổng quan 4 dự án
    # ----------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_bg(slide2)
    add_header(slide2, "SLIDE 2 / 10", "TỔNG QUAN 4 HƯỚNG TIẾP CẬN MÃ NGUỒN MỞ")

    rows, cols = 5, 4
    left, top, width, height = Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.0)
    table_shape = slide2.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(2.5)
    table.columns[1].width = Inches(2.3)
    table.columns[2].width = Inches(4.5)
    table.columns[3].width = Inches(2.433)

    headers = ["Dự án GitHub", "Công nghệ chính", "Hướng tiếp cận kỹ thuật", "Trạng thái Build/Test"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = COLOR_BLUE

    data = [
        ("1. FigmaToCode\n(bernaferrari/FigmaToCode)", "TypeScript, React 19, Turbo", "Client-side Plugin Sandbox, mã hóa Deterministic, xuất Zip Starter Project", "100% Vitest Passed"),
        ("2. Grida Design Engine\n(gridaco/grida)", "TypeScript, Rust, Kiwi Binary", "Multi-tier Design Engine, Dữ liệu trung gian IR (Intermediate Representation)", "Parsed & Checked Out"),
        ("3. Figma to React\n(kazuyaseki/figma-to-react)", "React, TS, Webpack 4", "Minimalist Figma Plugin, Mô hình buildTagTree tách biệt CSS", "9/9 Jest Passed, Build OK"),
        ("4. Figma Context MCP\n(GLips/Figma-Context-MCP)", "TypeScript, MCP Protocol", "Model Context Protocol Server, cung cấp Visual Specs JSON cho AI (LLM)", "249/250 Vitest, Build OK")
    ]

    for r_idx, row in enumerate(data):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_INNER
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.color.rgb = COLOR_GREEN if c_idx == 3 else TEXT_WHITE

    # ----------------------------------------------------
    # SLIDE 3: FigmaToCode
    # ----------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_bg(slide3)
    add_header(slide3, "SLIDE 3 / 10", "1. FIGMA TO CODE (DETERMINISTIC PLUGIN)")

    # Card 1
    c1 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.0))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = BORDER_COLOR
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "⚙️ Thuật toán Node Normalization (altNodes)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE

    bullets1 = [
        "Chuyển cây SceneNode phức tạp của Figma sang AltNode tinh gọn.",
        "Tự động phát hiện và xóa bỏ các Frame thừa không chứa style.",
        "Tự động gom nhóm thuộc tính padding (top/bottom/left/right) & gap.",
        "Tự động gom màu trùng thành CSS Variables & render SVG Icon fallback."
    ]
    for b in bullets1:
        p = tf1.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(10)

    # Card 2
    c2 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_COLOR
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "📦 Xuất Starter Project (.ZIP)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PURPLE

    bullets2 = [
        "Chạy thuần túy 100% trong Sandbox của Figma Plugin (Zero API/AI ngoài).",
        "Tích hợp thư viện fflate nén toàn bộ mã nguồn thành file .zip.",
        "Hỗ trợ đóng gói dự án Starter hoàn chỉnh: HTML Static, Tailwind CSS, React, Flutter, SwiftUI.",
        "Mã hóa Deterministic: Kết quả cố định, nhất quán 100% giữa các lần chạy."
    ]
    for b in bullets2:
        p = tf2.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(10)

    # ----------------------------------------------------
    # SLIDE 4: Grida
    # ----------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_bg(slide4)
    add_header(slide4, "SLIDE 4 / 10", "2. GRIDA DESIGN TO CODE (MULTI-TIER ENGINE)")

    c1 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.0))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = BORDER_COLOR
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "🧬 Đọc File Nhị Phân Kiwi (.fig)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN

    bullets1 = [
        "Tích hợp bộ parser nhị phân fig-kiwi chuyên dụng.",
        "Đọc trực tiếp file .fig / .deck offline mà không bắt buộc phải qua Figma REST API.",
        "Cấu trúc dữ liệu trung gian (@design-sdk / Grida IR Schema) hoàn toàn độc lập với framework đầu ra."
    ]
    for b in bullets1:
        p = tf1.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(12)

    c2 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_COLOR
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "📐 Suy Luận Flexbox / Grid Tự Động"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PURPLE

    bullets2 = [
        "Giải quyết bài toán các phần tử xếp vị trí tuyệt đối (Absolute positioning).",
        "Phân tích tọa độ hình học (x, y, w, h) để suy luận phân nhóm phần tử.",
        "Tự động chuyển đổi cấu trúc phức tạp thành HBox (Row), VBox (Column) hoặc Grid đáp ứng (Responsive Layout)."
    ]
    for b in bullets2:
        p = tf2.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(12)

    # ----------------------------------------------------
    # SLIDE 5: Figma to React
    # ----------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_bg(slide5)
    add_header(slide5, "SLIDE 5 / 10", "3. FIGMA TO REACT (TAG TREE ARCHITECTURE)")

    c1 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.0))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = BORDER_COLOR
    tf1 = c1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "🌳 Mô hình Tag Tree (buildTagTree.ts)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE

    p = tf1.add_paragraph()
    p.text = "Mẫu chuẩn mực nhất cho Figma Plugin: Bóc tách node Figma thành cây đối tượng Tag trung gian trước khi render JSX."
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(8)

    p = tf1.add_paragraph()
    p.text = "export type Tag = {\n  name: string;\n  isText: boolean;\n  textCharacters: string | null;\n  isImg: boolean;\n  properties: Property[];\n  css: CSSData;\n  children: Tag[];\n  node: SceneNode;\n}"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_GREEN
    p.font.name = 'Courier New'
    p.space_before = Pt(12)

    p = tf1.add_paragraph()
    p.text = "• getCssDataForTag.ts: Quy đổi trực tiếp các thuộc tính Figma sang CSS Rules một cách sạch rành mạch (flex-direction, padding, gap, border-radius).\n• buildCode.ts: Duyệt cây Tag đệ quy xuất ra React JSX kèm CSS Modules hoặc Styled-Components."
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(12)

    # ----------------------------------------------------
    # SLIDE 6: Figma Context MCP
    # ----------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_bg(slide6)
    add_header(slide6, "SLIDE 6 / 10", "4. FIGMA CONTEXT MCP (AI-NATIVE PROTOCOL)")

    c1 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.0))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = BORDER_COLOR
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "🤖 Model Context Protocol (MCP)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_ROSE

    bullets1 = [
        "Xu hướng công nghệ mới: Cầu nối dữ liệu giữa Figma API & AI LLM (Claude, Cursor).",
        "Không sinh code tĩnh cố định mà trích xuất Visual Specs / Design Tokens dạng JSON/Markdown.",
        "Tối ưu dung lượng token để AI đọc hiểu cấu trúc giao diện nhanh nhất."
    ]
    for b in bullets1:
        p = tf1.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(12)

    c2 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_COLOR
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🛠️ Các công cụ MCP cung cấp"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE

    bullets2 = [
        "get_figma_file: Lấy thông số tổng quan toàn bộ file thiết kế.",
        "get_figma_node: Lấy thông số chi tiết của 1 Component cụ thể.",
        "download_figma_images: Tự động tải hình ảnh assets từ Figma.",
        "Để AI tự biên dịch linh hoạt ra code tương thích với context dự án hiện tại."
    ]
    for b in bullets2:
        p = tf2.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(10)

    # ----------------------------------------------------
    # SLIDE 7: Bảng so sánh Ưu & Nhược điểm
    # ----------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_bg(slide7)
    add_header(slide7, "SLIDE 7 / 10", "BẢNG SO SÁNH ƯU & NHƯỢC ĐIỂM (PROS & CONS)")

    rows, cols = 5, 4
    left, top, width, height = Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.0)
    table_shape = slide7.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(3.8)
    table.columns[2].width = Inches(3.2)
    table.columns[3].width = Inches(2.533)

    headers = ["Dự án", "Ưu điểm (Pros)", "Nhược điểm (Cons)", "Khai thác cho Layoutly"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = COLOR_BLUE

    data7 = [
        ("FigmaToCode", "• Chạy 100% offline sandbox\n• Thuật toán gom padding/gap cực sạch\n• Xuất file ZIP dự án mẫu", "• Phụ thuộc vào thiết kế Auto Layout của Figma", "Học thuật toán gom nhóm CSS & xuất Starter Project ZIP"),
        ("Grida", "• Xử lý tốt layout vị trí tuyệt đối (Absolute)\n• Đọc được file nhị phân .fig", "• Kiến trúc rất phức tạp, mã nguồn lớn", "Học thiết kế lớp Dữ liệu Trung gian (IR)"),
        ("Figma to React", "• Code cực kỳ ngắn gọn, dễ hiểu\n• Tách biệt rõ giữa Tag Tree & CSS rules", "• Chưa hỗ trợ các ca layout quá phức tạp", "Làm mẫu cho Figma Plugin Bridge (layoutly-figma-bridge)"),
        ("Figma Context MCP", "• Đón đầu xu hướng AI Assistant\n• Code sinh ra linh hoạt theo prompt", "• Phụ thuộc vào AI LLM (tốn chi phí API)", "Định hướng mở rộng tính năng AI Assistant tương lai")
    ]

    for r_idx, row in enumerate(data7):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_INNER
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.color.rgb = COLOR_GREEN if c_idx == 1 else (COLOR_ROSE if c_idx == 2 else TEXT_WHITE)

    # ----------------------------------------------------
    # SLIDE 8: Bài học cho Layoutly
    # ----------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    add_bg(slide8)
    add_header(slide8, "SLIDE 8 / 10", "RÚT RA BÀI HỌC KỸ THUẬT CHO LAYOUTLY")

    coords = [
        (Inches(0.8), Inches(1.6), COLOR_BLUE, "1. Tầng Backend Parser (FigmaParserService.java)", "Áp dụng thuật toán altNodes của FigmaToCode: Thêm bước tiền xử lý lọc node visible=false, gom nhóm padding trùng lặp và phát hiện Icon SVG tự động."),
        (Inches(6.8), Inches(1.6), COLOR_PURPLE, "2. Tầng Cấu Trúc Trung Gian (Layoutly IR Schema)", "Học tập mô hình Tag Tree và Grida IR: Tạo lớp DTO LayoutTagDTO trong Spring Boot đứng giữa Parser Service và Code Generator Service."),
        (Inches(0.8), Inches(4.3), COLOR_GREEN, "3. Tầng Figma Plugin (layoutly-figma-bridge)", "Áp dụng mô hình truyền dữ liệu Message Passing tinh gọn từ Figma to React để trích xuất SceneNode từ Canvas nhanh chóng."),
        (Inches(6.8), Inches(4.3), COLOR_ROSE, "4. Đóng gói Dự án Starter (ExportService.java)", "Áp dụng cơ chế nén file ZIP từ FigmaToCode để đóng gói toàn bộ mã nguồn React/Vite/HTML cho người dùng tải về.")
    ]

    for x, y, col, header, text in coords:
        box = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.7), Inches(2.4))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = BORDER_COLOR
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = header
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = col
        p2 = tf.add_paragraph()
        p2.text = text
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_WHITE
        p2.space_before = Pt(8)

    # ----------------------------------------------------
    # SLIDE 9: Sơ đồ kiến trúc Layoutly
    # ----------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    add_bg(slide9)
    add_header(slide9, "SLIDE 9 / 10", "KIẾN TRÚC ĐỀ XUẤT CHO HỆ THỐNG LAYOUTLY")

    box9 = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.0))
    box9.fill.solid()
    box9.fill.fore_color.rgb = CARD_BG
    box9.line.color.rgb = BORDER_COLOR
    tf9 = box9.text_frame
    tf9.word_wrap = True

    p = tf9.paragraphs[0]
    p.text = "🔄 LUỒNG XỬ LÝ DỮ LIỆU TOÀN DIỆN CỦA LAYOUTLY"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PURPLE
    p.alignment = PP_ALIGN.CENTER

    steps = [
        ("1. Figma Plugin Bridge", "Trích xuất SceneNode thô từ Canvas và gửi Payload JSON về Backend API.", COLOR_BLUE),
        ("2. Spring Boot Backend", "FigmaParserService chạy altNodes Normalization & tạo LayoutTagDTO (IR).", COLOR_PURPLE),
        ("3. Code Generators", "CssGeneratorService & ReactGeneratorService sinh mã JSX + CSS Modules.", COLOR_GREEN),
        ("4. React FE & Export ZIP", "Hiển thị Live Preview trên FE Editor & nén Starter Project ZIP tải về.", COLOR_ROSE)
    ]

    for idx, (stitle, sdesc, scol) in enumerate(steps):
        sy = Inches(2.5 + idx * 1.0)
        sbox = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), sy, Inches(10.333), Inches(0.85))
        sbox.fill.solid()
        sbox.fill.fore_color.rgb = CARD_INNER
        sbox.line.color.rgb = scol
        stf = sbox.text_frame
        stf.word_wrap = True
        p = stf.paragraphs[0]
        p.text = f"{stitle}: "
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = scol
        run = p.add_run()
        run.text = sdesc
        run.font.size = Pt(12)
        run.font.bold = False
        run.font.color.rgb = TEXT_WHITE

    # ----------------------------------------------------
    # SLIDE 10: Kết luận
    # ----------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    add_bg(slide10)
    add_header(slide10, "SLIDE 10 / 10", "KẾT LUẬN & KẾ HOẠCH THỰC THI (NEXT STEPS)")

    # Card Top
    ct = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(1.8))
    ct.fill.solid()
    ct.fill.fore_color.rgb = CARD_BG
    ct.line.color.rgb = BORDER_COLOR
    tft = ct.text_frame
    tft.word_wrap = True
    p = tft.paragraphs[0]
    p.text = "📌 Kết luận Chuyên môn"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN

    p2 = tft.add_paragraph()
    p2.text = "Không có thuật toán đơn lẻ nào xử lý hoàn hảo 100% bài toán Figma-to-Code. Mô hình tối ưu nhất cho Layoutly là sự kết hợp giữa: Thuật toán làm sạch node của FigmaToCode + Kiến trúc IR trung gian của Grida + Cấu trúc Plugin gọn nhẹ của Figma to React."
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(6)

    # Card Bottom Left
    cb1 = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.6), Inches(5.7), Inches(3.0))
    cb1.fill.solid()
    cb1.fill.fore_color.rgb = CARD_BG
    cb1.line.color.rgb = BORDER_COLOR
    tfb1 = cb1.text_frame
    tfb1.word_wrap = True
    p = tfb1.paragraphs[0]
    p.text = "🗓️ Kế hoạch Hoàn thiện Layoutly"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE

    plans = [
        "Tuần 1: Nâng cấp FigmaParserService.java theo altNodes.",
        "Tuần 2: Chuẩn hóa lớp LayoutTagDTO (IR Schema) trong Backend.",
        "Tuần 3: Hoàn thiện ReactGeneratorService.java xuất JSX & ZIP.",
        "Tuần 4: Kiểm thử end-to-end & hoàn thiện báo cáo đồ án."
    ]
    for pl in plans:
        p = tfb1.add_paragraph()
        p.text = "• " + pl
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(6)

    # Card Bottom Right
    cb2 = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(3.6), Inches(5.7), Inches(3.0))
    cb2.fill.solid()
    cb2.fill.fore_color.rgb = CARD_BG
    cb2.line.color.rgb = BORDER_COLOR
    tfb2 = cb2.text_frame
    tfb2.word_wrap = True
    p = tfb2.paragraphs[0]
    p.text = "\nCẢM ƠN THẦY CÔ & HỘI ĐỒNG!"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    p.alignment = PP_ALIGN.CENTER

    p2 = tfb2.add_paragraph()
    p2.text = "\nSẵn sàng nhận câu hỏi phản biện & đóng góp ý kiến."
    p2.font.size = Pt(14)
    p2.font.color.rgb = TEXT_MUTED
    p2.alignment = PP_ALIGN.CENTER

    # Save output
    output_path = "c:\\Users\\thanh\\Desktop\\Layoutly\\layoutly_presentation_slides.pptx"
    prs.save(output_path)
    print(f"SUCCESS: Presentation saved to {output_path}")

if __name__ == "__main__":
    create_presentation()
